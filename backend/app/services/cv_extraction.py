"""Read-only digital extraction. No OCR, persistence, candidate parsing or network."""
from pathlib import Path
import re
from threading import Lock
import unicodedata
from uuid import UUID

import pymupdf
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.schemas.cv_extraction import (
    MAX_SOURCE_UNITS, MAX_TEXT_CHARS, DocxTextBlock, ExtractedCvDocument, PdfPageText,
)
from app.schemas.cv_upload import CvBatchInput
from app.services.cv_upload import MAX_BATCH_BYTES, MAX_FILES, MAX_FILE_BYTES, CvUploadService, validate_document

MAX_PDF_PAGES = 200
MIN_PDF_ALNUM = 80
MIN_PAGE_ALNUM = 20
MIN_MEANINGFUL_PAGE_RATIO = 0.5
STORED_NAME = re.compile(r"^([0-9a-f]{8}-[0-9a-f]{4}-4[0-9a-f]{3}-[89ab][0-9a-f]{3}-[0-9a-f]{12})\.(pdf|docx)$")
# MuPDF is not safe for concurrent threaded use. Serialize PDF access in this service.
PDF_LOCK = Lock()
PUBLIC_ERRORS = {
    "batch_not_found": "The uploaded batch was not found.",
    "unsafe_batch": "The stored batch cannot be processed safely.",
    "extraction_unavailable": "CV text extraction could not be completed.",
}


class CvExtractionError(Exception):
    def __init__(self, code: str, status: int = 400):
        self.code, self.status = code, status
        super().__init__(PUBLIC_ERRORS[code])


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return "".join(c for c in text if c in "\n\t" or unicodedata.category(c) not in {"Cc", "Cf", "Cs"}).strip()


def alnum_count(text: str) -> int:
    return sum(c.isalnum() for c in text)


def docx_blocks(container, in_cell=False, depth=0):
    """Walk paragraphs and tables in XML order; visit each merged cell once."""
    if depth > 16:
        raise ValueError("Table nesting limit")
    for block in container.iter_inner_content():
        if isinstance(block, Paragraph):
            text = normalize_text(block.text)
            if text:
                yield ("table_cell" if in_cell else "paragraph", text)
        elif isinstance(block, Table):
            seen = set()
            for row in block.rows:
                for cell in row.cells:
                    if cell._tc not in seen:
                        seen.add(cell._tc)
                        yield from docx_blocks(cell, in_cell=True, depth=depth + 1)


class CvExtractionService:
    def __init__(self, storage_root: Path):
        self.root = CvUploadService(storage_root).root

    def extract_batch(self, job_reference: str, batch_id: UUID) -> tuple[ExtractedCvDocument, ...]:
        try:
            reference = CvBatchInput(job_reference=job_reference).job_reference
            batch_id = UUID(str(batch_id))
        except ValueError:
            raise CvExtractionError("unsafe_batch") from None
        job = self.root / reference
        batch = job / str(batch_id)
        if job.is_symlink() or batch.is_symlink() or job.is_junction() or batch.is_junction():
            raise CvExtractionError("unsafe_batch")
        if job.exists() and not job.is_dir():
            raise CvExtractionError("unsafe_batch")
        if not batch.exists():
            raise CvExtractionError("batch_not_found", 404)
        if not batch.is_dir() or job.resolve().parent != self.root or batch.resolve().parent != job.resolve():
            raise CvExtractionError("unsafe_batch")
        files = sorted(batch.iterdir(), key=lambda path: path.name)
        if not 1 <= len(files) <= MAX_FILES:
            raise CvExtractionError("unsafe_batch")
        total = 0
        for path in files:
            if (not STORED_NAME.fullmatch(path.name) or path.is_symlink() or path.is_junction()
                or not path.is_file() or path.resolve().parent != batch.resolve()):
                raise CvExtractionError("unsafe_batch")
            size = path.stat().st_size
            total += size
            if not 0 < size <= MAX_FILE_BYTES or total > MAX_BATCH_BYTES:
                raise CvExtractionError("unsafe_batch")
        return tuple(self._extract_file(path) for path in files)

    def _extract_file(self, path: Path) -> ExtractedCvDocument:
        upload_id, file_type = UUID(path.stem), path.suffix[1:]
        try:
            validate_document(path, file_type)
            if file_type == "pdf":
                units = self._pdf_units(path)
                meaningful = sum(alnum_count(unit.text) >= MIN_PAGE_ALNUM for unit in units)
                sufficient = (sum(alnum_count(unit.text) for unit in units) >= MIN_PDF_ALNUM
                    and meaningful / len(units) >= MIN_MEANINGFUL_PAGE_RATIO)
                status = "text_extracted" if sufficient else "ocr_required"
            else:
                units = []
                chars = 0
                for kind, text in docx_blocks(Document(path)):
                    chars += len(text) + (2 if units else 0)
                    if chars > MAX_TEXT_CHARS or len(units) >= MAX_SOURCE_UNITS:
                        raise ValueError("Extraction limit")
                    units.append(DocxTextBlock(block_index=len(units) + 1, block_type=kind, text=text))
                status = "text_extracted" if any(alnum_count(unit.text) for unit in units) else "extraction_failed"
            text = "\n\n".join(unit.text for unit in units)
            return ExtractedCvDocument(upload_id=upload_id, file_type=file_type,
                extraction_status=status, text=text, text_char_count=len(text), source_units=tuple(units))
        except Exception:
            # Parser details and candidate text must never cross the public boundary.
            return ExtractedCvDocument(upload_id=upload_id, file_type=file_type,
                extraction_status="extraction_failed", text="", text_char_count=0, source_units=())

    @staticmethod
    def _pdf_units(path: Path) -> list[PdfPageText]:
        with PDF_LOCK:
            errors = pymupdf.TOOLS.mupdf_display_errors()
            warnings = pymupdf.TOOLS.mupdf_display_warnings()
            pymupdf.TOOLS.mupdf_display_errors(False)
            pymupdf.TOOLS.mupdf_display_warnings(False)
            try:
                with pymupdf.open(path) as document:
                    if document.needs_pass or not 1 <= document.page_count <= MAX_PDF_PAGES:
                        raise ValueError("Unsupported PDF")
                    units, chars = [], 0
                    for index, page in enumerate(document):
                        text = normalize_text(page.get_text("text", sort=True))
                        chars += len(text) + (2 if units else 0)
                        if chars > MAX_TEXT_CHARS:
                            raise ValueError("Extraction limit")
                        units.append(PdfPageText(page_number=index + 1, text=text))
                    return units
            finally:
                pymupdf.TOOLS.mupdf_display_errors(errors)
                pymupdf.TOOLS.mupdf_display_warnings(warnings)
                pymupdf.TOOLS.mupdf_warnings(reset=True)
