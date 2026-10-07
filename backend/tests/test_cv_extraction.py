"""Synthetic documents only. Exercise actual parsers and public privacy boundaries."""
import asyncio
from io import BytesIO
from uuid import uuid4
from zipfile import ZipFile

import httpx
import pymupdf
import pytest
from docx import Document

from app.api.routers import cv_uploads as route
from app.core.config import Settings
from app.main import app
from app.schemas.cv_extraction import DocxTextBlock, PdfPageText
from app.services import cv_extraction as service

REFERENCE = "TT-JOB-30BB2CDF9A5CCCA9"
TEXT = "SYNTHETIC ONLY NOT A REAL CV. Python backend development, automated testing, REST APIs and relational database projects."


def pdf_bytes(texts=(TEXT,)):
    with pymupdf.open() as document:
        for text in texts:
            page = document.new_page()
            if text:
                page.insert_textbox(pymupdf.Rect(40, 40, 550, 750), text, fontsize=11)
        return document.tobytes()


def docx_bytes(empty=False):
    document = Document()
    if not empty:
        document.add_paragraph("SYNTHETIC first paragraph")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "SYNTHETIC first cell"
        table.cell(0, 1).text = "SYNTHETIC second cell"
        document.add_paragraph("SYNTHETIC last paragraph")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.fixture(autouse=True)
def storage(monkeypatch, tmp_path):
    root = tmp_path / "private"
    monkeypatch.setattr(route, "settings", Settings(_env_file=None, cv_upload_root=root))
    async def forbidden(*args, **kwargs):
        pytest.fail("Extraction must not contact external services.")
    monkeypatch.setattr(httpx.AsyncHTTPTransport, "handle_async_request", forbidden)
    return root


def stored(root, *files):
    batch_id = uuid4()
    batch = root / REFERENCE / str(batch_id)
    batch.mkdir(parents=True)
    paths = []
    for kind, content in files:
        path = batch / f"{uuid4()}.{kind}"
        path.write_bytes(content)
        paths.append(path)
    return batch_id, batch, paths


def request(batch_id, reference=REFERENCE):
    async def send():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
            return await client.post(f"/api/v1/cv-uploads/{reference}/{batch_id}/extract")
    return asyncio.run(send())


@pytest.mark.parametrize("texts", [(TEXT,), ("First page " + TEXT, "Second page " + TEXT)])
def test_pdf_actual_text_page_order_and_originals(storage, texts):
    batch_id, _, paths = stored(storage, ("pdf", pdf_bytes(texts)))
    original = paths[0].read_bytes()
    result, = service.CvExtractionService(storage).extract_batch(REFERENCE, batch_id)
    assert result.extraction_status == "text_extracted"
    assert [unit.page_number for unit in result.source_units] == list(range(1, len(texts) + 1))
    assert all(isinstance(unit, PdfPageText) for unit in result.source_units)
    for unit, expected in zip(result.source_units, texts):
        assert " ".join(unit.text.split()) == expected
    assert result.text_char_count == len(result.text)
    assert paths[0].read_bytes() == original
    assert TEXT not in repr(result)
    public = request(batch_id)
    assert public.status_code == 200
    assert public.json()["files"][0]["source_unit_count"] == len(texts)
    for private in ["SYNTHETIC", str(storage), str(paths[0]), "page_number", '"text"', "filename"]:
        assert private not in public.text


@pytest.mark.parametrize("texts", [("",), ("Tiny text",), ("!!!" * 60,), (TEXT, "", "")])
def test_insufficient_pdf_text_routes_to_ocr_without_doing_ocr(storage, texts, monkeypatch):
    monkeypatch.setattr(pymupdf.Page, "get_textpage_ocr", lambda *a, **k: pytest.fail("OCR is not implemented"))
    batch_id, _, _ = stored(storage, ("pdf", pdf_bytes(texts)))
    result = request(batch_id).json()["files"][0]
    assert result["extraction_status"] == "ocr_required"
    assert result["ocr_recommended"] is True
    assert result["source_unit_count"] == len(texts)


@pytest.mark.parametrize("texts,status", [(('A' * 79,), 'ocr_required'), (('A' * 80,), 'text_extracted'),
    ((TEXT, ''), 'text_extracted'), ((TEXT, '', ''), 'ocr_required')])
def test_explicit_ocr_threshold_boundaries(storage, texts, status):
    batch_id, _, _ = stored(storage, ("pdf", pdf_bytes(texts)))
    result, = service.CvExtractionService(storage).extract_batch(REFERENCE, batch_id)
    assert result.extraction_status == status


def test_docx_preserves_paragraph_table_interleaving_and_original(storage):
    batch_id, _, paths = stored(storage, ("docx", docx_bytes()))
    original = paths[0].read_bytes()
    result, = service.CvExtractionService(storage).extract_batch(REFERENCE, batch_id)
    assert result.extraction_status == "text_extracted"
    assert all(isinstance(unit, DocxTextBlock) for unit in result.source_units)
    assert [unit.block_index for unit in result.source_units] == [1, 2, 3, 4]
    assert [unit.block_type for unit in result.source_units] == ["paragraph", "table_cell", "table_cell", "paragraph"]
    assert [unit.text for unit in result.source_units] == ["SYNTHETIC first paragraph", "SYNTHETIC first cell", "SYNTHETIC second cell", "SYNTHETIC last paragraph"]
    assert paths[0].read_bytes() == original
    response = request(batch_id)
    assert response.status_code == 200
    assert response.json()["files"][0]["source_unit_count"] == 4
    assert "SYNTHETIC" not in response.text


def test_merged_and_nested_docx_tables_are_not_duplicated(storage):
    document = Document()
    table = document.add_table(rows=1, cols=2)
    cell = table.cell(0, 0).merge(table.cell(0, 1))
    cell.text = "Outer synthetic"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "Inner synthetic"
    cell.add_paragraph("After nested synthetic")
    output = BytesIO(); document.save(output)
    batch_id, _, _ = stored(storage, ("docx", output.getvalue()))
    result, = service.CvExtractionService(storage).extract_batch(REFERENCE, batch_id)
    assert [unit.text for unit in result.source_units] == ["Outer synthetic", "Inner synthetic", "After nested synthetic"]


@pytest.mark.parametrize("kind,content", [("docx", b"PK\x03\x04BROKEN"), ("pdf", b"%PDF-1.4\nPRIVATE_PARSER_DETAILS\n%%EOF"), ("docx", docx_bytes(empty=True))], ids=["malformed-docx", "malformed-pdf", "empty-docx"])
def test_malformed_or_empty_document_safe_failure(storage, kind, content, capsys, caplog):
    batch_id, _, paths = stored(storage, (kind, content))
    response = request(batch_id)
    assert response.status_code == 200
    item = response.json()["files"][0]
    assert item["extraction_status"] == "extraction_failed"
    assert item["ocr_recommended"] is False
    assert paths[0].read_bytes() == content
    output = capsys.readouterr()
    for private in ["PRIVATE_PARSER_DETAILS", str(storage)]:
        assert private not in response.text + output.out + output.err + caplog.text


def test_mixed_batch_deterministic_order_is_non_destructive(storage):
    batch_id, batch, paths = stored(storage, ("pdf", pdf_bytes()), ("pdf", pdf_bytes(("",))),
        ("docx", docx_bytes()), ("pdf", b"%PDF-1.4\nbroken\n%%EOF"))
    originals = {path.name: path.read_bytes() for path in paths}
    first = request(batch_id); second = request(batch_id)
    assert first.status_code == 200
    assert first.json() == second.json()
    assert first.json()["processed_count"] == 4
    assert [item["upload_id"] for item in first.json()["files"]] == [path.stem for path in sorted(paths)]
    assert {item["extraction_status"] for item in first.json()["files"]} == {"text_extracted", "ocr_required", "extraction_failed"}
    assert {path.name: path.read_bytes() for path in batch.iterdir()} == originals


def test_docx_broken_xml_passes_package_checks_but_parser_fails_safely(storage):
    output = BytesIO()
    with ZipFile(BytesIO(docx_bytes())) as source, ZipFile(output, "w") as target:
        for name in source.namelist():
            target.writestr(name, b"<BROKEN_PRIVATE_XML" if name == "word/document.xml" else source.read(name))
    batch_id, _, paths = stored(storage, ("docx", output.getvalue()))
    service.validate_document(paths[0], "docx")
    response = request(batch_id)
    assert response.status_code == 200
    assert response.json()["files"][0]["extraction_status"] == "extraction_failed"
    assert "BROKEN_PRIVATE_XML" not in response.text


@pytest.mark.parametrize("reference,batch", [("not-a-reference", str(uuid4())), (REFERENCE, "not-a-uuid"), ("TT-JOB-" + "X" * 16, str(uuid4()))])
def test_invalid_identifiers_are_sanitized_422(reference, batch):
    response = request(batch, reference)
    assert response.status_code == 422
    assert reference not in response.text
    assert batch not in response.text


def test_missing_batch_404(storage):
    response = request(uuid4())
    assert response.status_code == 404
    assert response.json()["code"] == "batch_not_found"
    assert str(storage) not in response.text


def test_job_directory_cannot_be_a_regular_file(storage):
    storage.mkdir()
    (storage / REFERENCE).write_bytes(b"synthetic unexpected layout")
    assert request(uuid4()).status_code == 400


@pytest.mark.parametrize("layout", ["unexpected_name", "directory", "empty", "oversized", "too_many"])
def test_unsafe_layout_rejected_before_parsing(storage, layout, monkeypatch):
    batch_id, batch, paths = stored(storage, ("pdf", pdf_bytes()))
    if layout == "unexpected_name": paths[0].rename(batch / "original-name.pdf")
    elif layout == "directory": (batch / f"{uuid4()}.docx").mkdir()
    elif layout == "empty": paths[0].unlink()
    elif layout == "oversized": monkeypatch.setattr(service, "MAX_FILE_BYTES", 1)
    else:
        for _ in range(20): (batch / f"{uuid4()}.pdf").write_bytes(b"x")
    monkeypatch.setattr(service.CvExtractionService, "_extract_file", lambda *a: pytest.fail("Unsafe layout must not be parsed"))
    response = request(batch_id)
    assert response.status_code == 400
    assert response.json()["code"] == "unsafe_batch"
    assert str(storage) not in response.text


@pytest.mark.parametrize("target", ["batch", "job", "file"])
def test_symlink_layout_rejected(storage, tmp_path, target):
    batch_id, batch, paths = stored(storage, ("pdf", pdf_bytes()))
    original = batch if target == "batch" else batch.parent if target == "job" else paths[0]
    moved = tmp_path / f"outside-{target}"
    original.rename(moved)
    try:
        original.symlink_to(moved, target_is_directory=target != "file")
    except OSError:
        pytest.skip("Creating symlinks requires OS permission on this host.")
    response = request(batch_id)
    assert response.status_code == 400
    assert response.json()["code"] == "unsafe_batch"


def test_sanitized_service_failure(storage, monkeypatch):
    monkeypatch.setattr(service.CvExtractionService, "extract_batch", lambda *a: (_ for _ in ()).throw(RuntimeError("PRIVATE_PATH_SECRET")))
    response = request(uuid4())
    assert response.status_code == 500
    assert response.json()["code"] == "extraction_unavailable"
    assert "PRIVATE" not in response.text


@pytest.mark.parametrize("target", ["batch", "job", "file"])
def test_detected_symlink_layout_is_rejected_on_all_platforms(storage, monkeypatch, target):
    # Windows may deny creation of real symlinks; still exercise each guard there.
    from pathlib import Path
    batch_id, batch, paths = stored(storage, ("pdf", pdf_bytes()))
    selected = batch if target == "batch" else batch.parent if target == "job" else paths[0]
    original = Path.is_symlink
    monkeypatch.setattr(Path, "is_symlink", lambda path: path == selected or original(path))
    assert request(batch_id).status_code == 400


@pytest.mark.parametrize("limit", ["MAX_TEXT_CHARS", "MAX_PDF_PAGES", "MAX_SOURCE_UNITS"])
def test_resource_limits_fail_per_file(storage, monkeypatch, limit):
    kind = "docx" if limit == "MAX_SOURCE_UNITS" else "pdf"
    batch_id, _, _ = stored(storage, (kind, docx_bytes() if kind == "docx" else pdf_bytes()))
    monkeypatch.setattr(service, limit, 0)
    assert request(batch_id).json()["files"][0]["extraction_status"] == "extraction_failed"


def test_normalization_removes_unsafe_controls_without_semantic_rewrite():
    assert service.normalize_text("  Python\r\nAPI\x00\x01\u202e\tSkills  ") == "Python\nAPI\tSkills"


def test_extraction_route_documented_without_private_text_contract():
    operation = app.openapi()["paths"]["/api/v1/cv-uploads/{job_reference}/{batch_id}/extract"]["post"]
    assert operation["responses"]["200"]["content"]["application/json"]["schema"]["$ref"].endswith("CvExtractionResponse")
    assert "ExtractedCvDocument" not in app.openapi()["components"]["schemas"]
