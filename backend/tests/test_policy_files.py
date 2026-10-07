import io
import struct
import zlib
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from app.rag import policy_files
from app.rag.policy_files import PolicyFileError, extract_policy_text


def write_pdf(path, pages, password=None, compress=False):
    """Real small text PDF fixture without a third document-generation library."""
    writer = PdfWriter()
    for text in pages:
        page = writer.add_blank_page(width=612, height=792)
        if text is not None:
            font = DictionaryObject({
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            })
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})
            })
            escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 50 750 Td ({escaped}) Tj ET".encode("ascii"))
            page[NameObject("/Contents")] = writer._add_object(stream.flate_encode() if compress else stream)
    if password is not None:
        writer.encrypt(password)
    writer.write(path)


def png_bytes():
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(b"\x00\xff\xff\xff")) + chunk(b"IEND", b""))


@pytest.mark.parametrize("extension", [".txt", ".md", ".TXT", ".MD"])
def test_utf8_text_preserves_headings_normalizes_endings_and_whitespace(tmp_path, extension):
    path = tmp_path / ("policy" + extension)
    path.write_bytes(b"\xef\xbb\xbf# Privacy\r\n\r\nProtect synthetic information.  \r\n")
    assert extract_policy_text(path) == "# Privacy\n\nProtect synthetic information."


@pytest.mark.parametrize("case", ["unsupported", "directory", "oversized", "empty", "missing", "encoding", "blank", "too_much_text"])
def test_invalid_files_fail_without_path_or_body(tmp_path, case):
    path = tmp_path / "PRIVATE_PATH_POLICY.txt"
    if case == "unsupported":
        path = path.with_suffix(".doc")
        path.write_text("PRIVATE_BODY")
    elif case == "directory":
        path.mkdir()
    elif case == "oversized":
        with path.open("wb") as stream:
            stream.truncate(policy_files.MAX_FILE_BYTES + 1)
    elif case == "empty":
        path.touch()
    elif case == "encoding":
        path.write_bytes(b"PRIVATE_BODY\xff")
    elif case == "blank":
        path.write_text(" \n\t ")
    elif case == "too_much_text":
        path.write_text("x" * (policy_files.MAX_POLICY_TEXT_LENGTH + 1))
    with pytest.raises(PolicyFileError) as error:
        extract_policy_text(path)
    assert "PRIVATE" not in str(error.value)
    assert str(tmp_path) not in str(error.value)


def test_pdf_text_page_order_and_blank_page(tmp_path):
    path = tmp_path / "policy.pdf"
    write_pdf(path, ["First policy page", None, "Third policy page"])
    assert extract_policy_text(path) == "First policy page\n\nThird policy page"
    assert extract_policy_text(path) == extract_policy_text(path)


def test_pdf_none_text_is_safe(monkeypatch, tmp_path):
    import pypdf
    path = tmp_path / "policy.pdf"
    path.write_bytes(b"synthetic-test")

    class Page:
        def __init__(self, text):
            self.text = text
        def extract_text(self):
            return self.text

    class Reader:
        is_encrypted = False
        pages = [Page(None), Page("Safe policy body")]

    monkeypatch.setattr(pypdf, "PdfReader", lambda *args, **kwargs: Reader())
    assert extract_policy_text(path) == "Safe policy body"


def test_pdf_compressed_stream_expansion_is_bounded(tmp_path):
    path = tmp_path / "policy.pdf"
    write_pdf(path, ["x" * (policy_files.MAX_PDF_STREAM_BYTES + 1)], compress=True)
    assert path.stat().st_size < policy_files.MAX_FILE_BYTES
    with pytest.raises(PolicyFileError, match="PDF could not be read safely"):
        extract_policy_text(path)


@pytest.mark.parametrize("case", ["no_text", "malformed", "encrypted", "empty_password", "too_many_pages"])
def test_pdf_controlled_failure_and_credential_free_opening(tmp_path, case):
    path = tmp_path / "PRIVATE_PATH.pdf"
    if case == "malformed":
        path.write_bytes(b"PRIVATE_DOCUMENT_BODY invalid pdf")
    elif case == "no_text":
        write_pdf(path, [None])
    elif case == "too_many_pages":
        write_pdf(path, [None] * (policy_files.MAX_PDF_PAGES + 1))
    else:
        write_pdf(path, ["Synthetic policy body"], password="" if case == "empty_password" else "PRIVATE_PASSWORD")
    if case == "empty_password":
        assert extract_policy_text(path) == "Synthetic policy body"
    else:
        with pytest.raises(PolicyFileError) as error:
            extract_policy_text(path)
        assert "PRIVATE" not in str(error.value)
        if case == "no_text":
            assert "OCR" in str(error.value)


def test_docx_heading_paragraphs_table_order_and_merged_cells(tmp_path):
    document = Document()
    document.add_heading("Privacy", level=2)
    document.add_paragraph("First paragraph")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Left cell"
    table.cell(0, 1).text = "Right cell"
    document.add_paragraph("Last paragraph")
    merged = document.add_table(rows=1, cols=2)
    merged.cell(0, 0).merge(merged.cell(0, 1)).text = "Merged text"
    path = tmp_path / "policy.docx"
    document.save(path)
    assert extract_policy_text(path) == (
        "## Privacy\n\nFirst paragraph\n\nLeft cell | Right cell\n\nLast paragraph\n\nMerged text"
    )


@pytest.mark.parametrize("case", ["empty", "empty_table", "corrupt", "archive_expansion", "macros"])
def test_docx_invalid_content(tmp_path, case):
    path = tmp_path / "PRIVATE_PATH.docx"
    if case == "corrupt":
        path.write_bytes(b"PRIVATE_BODY not a docx")
    elif case in {"archive_expansion", "macros"}:
        with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
            archive.writestr("word/vbaProject.bin" if case == "macros" else "word/document.xml",
                             b"x" if case == "macros" else b"x" * (policy_files.MAX_DOCX_EXPANDED_BYTES + 1))
    else:
        document = Document()
        if case == "empty_table":
            document.add_table(rows=1, cols=2)
        document.save(path)
    with pytest.raises(PolicyFileError) as error:
        extract_policy_text(path)
    assert "PRIVATE" not in str(error.value)


def test_docx_embedded_image_is_ignored_without_network(tmp_path, monkeypatch):
    import socket
    monkeypatch.setattr(socket, "create_connection", lambda *args, **kwargs: pytest.fail("No network allowed"))
    document = Document()
    document.add_paragraph("Policy text before image")
    document.add_picture(io.BytesIO(png_bytes()))
    document.add_paragraph("Policy text after image")
    path = tmp_path / "policy.docx"
    document.save(path)
    assert extract_policy_text(path) == "Policy text before image\n\nPolicy text after image"
